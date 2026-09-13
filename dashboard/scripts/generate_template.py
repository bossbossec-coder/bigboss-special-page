"""各店舗が実際の数値を入力するための、記入用Excelテンプレートを生成するスクリプト。

生成物: dashboard/templates/日次売上_入力テンプレート.xlsx

- 「使い方」シート: 運用ルールの説明
- 「売上入力」シート: 実際に日々入力していくシート（店舗名はプルダウン選択）
- 「店舗マスタ」シート: プルダウンに出す店舗名の一覧（実際の店舗名に書き換える）

実行方法:
    python3 scripts/generate_template.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT_PATH = (
    Path(__file__).resolve().parent.parent / "templates" / "日次売上_入力テンプレート.xlsx"
)

HEADER_FILL = PatternFill(start_color="4E79A7", end_color="4E79A7", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
INPUT_ROWS = 800  # あらかじめ罫線・書式・プルダウンを用意しておく行数（2年分以上の余裕を持たせる）
STORE_ROWS = 30  # 店舗マスタに用意しておく行数

# 動作確認用のサンプル店舗名（実際の店舗名に書き換えてください）。
# 「外販」はスポット売上専用の特別な値で、店舗名として書き換えないでください
# （ダッシュボード側で店舗の合計・前年比などから常に除外され、別枠で集計されます）。
PLACEHOLDER_STORES = [
    "渋谷店", "新宿店", "池袋店", "横浜店", "大宮店",
    "千葉店", "立川店", "大阪店", "名古屋店", "福岡店",
    "外販",
]


def build_usage_sheet(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "使い方"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 90

    lines = [
        ("このファイルの使い方", True),
        ("", False),
        ("① 「店舗マスタ」シートの店舗名を、実際の店舗名に書き換えてください。", False),
        ("　（最初はサンプルの店舗名が入っています。行数が足りなければ追加してください）", False),
        ("", False),
        ("② 「売上入力」シートに、日々の実績を1行ずつ入力してください。", False),
        ("　・日付：その日の日付", False),
        ("　・店舗名：プルダウンから選択（店舗マスタシートの一覧から選べます）", False),
        ("　・売上金額：その日・その店舗の売上合計（数値のみ）", False),
        ("　・客数：来客数（任意・分かる場合のみ）", False),
        ("　・備考：任意メモ（臨時休業など）", False),
        ("", False),
        ("③ このファイルをダッシュボードが読み込む共有フォルダ（data/incoming）に保存してください。", False),
        ("　・店舗ごとに1ファイルを作り、日々その中に行を追記していく運用を想定しています。", False),
        ("　・ファイル名は例:  渋谷店_売上.xlsx  のように店舗が分かる名前にしてください。", False),
        ("", False),
        ("【注意】", True),
        ("・「店舗名」の表記は必ず店舗マスタの一覧と統一してください（表記ゆれがあると別店舗として集計されます）。", False),
        ("・売上金額には文字や単位（円など）を含めず、数値のみを入力してください。", False),
        ("・行の入力順序は関係ありません（新しい日付が上でも下でも、順不同でも問題なく集計されます）。", False),
        ("・日付は必ず「日付」として入力してください（例: 2026/9/10 や 2026-9-10 と入力すればExcelが自動で", False),
        ("　日付として認識し、表示は yyyy-mm-dd になります）。「2026年9月10日」のような文字列や、", False),
        ("　20260910のような数値としての入力はできません。", False),
        ("・過去1年以上のデータをまとめて入力しても問題ありません（前年同日比や着地予測の精度が上がります）。", False),
        (f"　あらかじめ{INPUT_ROWS}行分（約2年強）に書式とプルダウンを用意していますが、それを超えて入力する場合は、", False),
        ("　最後の行を選択してコピーし、続きの行に貼り付ければ書式・プルダウンを引き継げます。", False),
        ("", False),
        ("【外販（スポット売上）について】", True),
        ("・店舗名で「外販」を選ぶと、店舗の売上とは別枠の「外販売上」として記録されます。", False),
        ("・外販の実績は、店舗別の合計・ランキング・前年比などには一切含まれません", False),
        ("　（ダッシュボード下部に、外販だけの本日・前年同日比・当月累計が別途表示されます）。", False),
        ("・スポットでたまにしか売上が発生しない場合のみ使う想定です。日付・売上金額の入力方法は店舗と同じです。", False),
    ]

    row = 1
    for text, is_title in lines:
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(bold=True, size=14) if is_title and row == 1 else (
            Font(bold=True) if is_title else Font()
        )
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1


def build_store_master_sheet(wb: Workbook) -> str:
    ws = wb.create_sheet("店舗マスタ")
    ws.column_dimensions["A"].width = 24

    header = ws.cell(row=1, column=1, value="店舗名")
    header.font = HEADER_FONT
    header.fill = HEADER_FILL
    header.alignment = Alignment(horizontal="center")

    for i in range(STORE_ROWS):
        r = i + 2
        cell = ws.cell(row=r, column=1)
        cell.border = THIN_BORDER
        if i < len(PLACEHOLDER_STORES):
            cell.value = PLACEHOLDER_STORES[i]

    return f"店舗マスタ!$A$2:$A${STORE_ROWS + 1}"


def build_input_sheet(wb: Workbook, store_range: str) -> None:
    ws = wb.create_sheet("売上入力")
    headers = ["日付", "店舗名", "売上金額", "客数", "備考"]
    widths = [14, 16, 16, 10, 30]

    for col, (title, width) in enumerate(zip(headers, widths), start=1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = width
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"

    store_validation = DataValidation(
        type="list",
        formula1=store_range,
        allow_blank=True,
        showDropDown=False,  # openpyxl/Excelの仕様上、Falseにするとプルダウン矢印が表示される
        showErrorMessage=True,
    )
    store_validation.error = "店舗マスタの一覧にある店舗名を選択してください。"
    store_validation.errorTitle = "店舗名エラー"
    ws.add_data_validation(store_validation)

    for r in range(2, INPUT_ROWS + 2):
        date_cell = ws.cell(row=r, column=1)
        date_cell.number_format = "yyyy-mm-dd"
        date_cell.border = THIN_BORDER

        store_cell = ws.cell(row=r, column=2)
        store_cell.border = THIN_BORDER
        store_validation.add(store_cell)

        sales_cell = ws.cell(row=r, column=3)
        sales_cell.number_format = "#,##0"
        sales_cell.border = THIN_BORDER

        customers_cell = ws.cell(row=r, column=4)
        customers_cell.number_format = "#,##0"
        customers_cell.border = THIN_BORDER

        note_cell = ws.cell(row=r, column=5)
        note_cell.border = THIN_BORDER


def main() -> None:
    wb = Workbook()
    build_usage_sheet(wb)
    store_range = build_store_master_sheet(wb)
    build_input_sheet(wb, store_range)
    wb.active = wb["売上入力"]

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"テンプレートを作成しました: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
